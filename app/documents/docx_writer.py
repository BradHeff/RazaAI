from datetime import date
from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from .models import DocumentModel


NAVY = "173C5E"
TEAL = "2A7C91"
PALE_BLUE = "EDF5F8"
PALE_GREY = "F5F7F9"
MID_GREY = "D6DEE4"
MUTED = "66737D"
TIP_BG = "EAF4F7"
IMPORTANT_BG = "FFF5DB"
EXAMPLE_BG = "EDF7EF"


class DocxWriter:
    def __init__(self):
        self.NAVY = NAVY
        self.TEAL = TEAL
        self.MUTED = MUTED

    def _apply_theme(self, model):
        theme = (getattr(model, "metadata", {}) or {}).get("theme") or {}
        self.NAVY = theme.get("primary_color") or NAVY
        self.TEAL = theme.get("accent_color") or TEAL
        self.MUTED = theme.get("muted_color") or MUTED

    @staticmethod
    def _set_repeat_table_header(row):
        tr_pr = row._tr.get_or_add_trPr()
        tbl_header = OxmlElement("w:tblHeader")
        tbl_header.set(qn("w:val"), "true")
        tr_pr.append(tbl_header)

    @staticmethod
    def _shade_cell(cell, fill):
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), fill)
        tc_pr.append(shd)

    @staticmethod
    def _set_cell_margins(cell, top=90, start=90, bottom=90, end=90):
        tc_pr = cell._tc.get_or_add_tcPr()
        tc_mar = tc_pr.first_child_found_in("w:tcMar")
        if tc_mar is None:
            tc_mar = OxmlElement("w:tcMar")
            tc_pr.append(tc_mar)
        for margin, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
            node = tc_mar.find(qn(f"w:{margin}"))
            if node is None:
                node = OxmlElement(f"w:{margin}")
                tc_mar.append(node)
            node.set(qn("w:w"), str(value))
            node.set(qn("w:type"), "dxa")

    @staticmethod
    def _set_cell_border(cell, **edges):
        tc_pr = cell._tc.get_or_add_tcPr()
        borders = tc_pr.first_child_found_in("w:tcBorders")
        if borders is None:
            borders = OxmlElement("w:tcBorders")
            tc_pr.append(borders)
        for edge, attrs in edges.items():
            tag = f"w:{edge}"
            node = borders.find(qn(tag))
            if node is None:
                node = OxmlElement(tag)
                borders.append(node)
            for key, value in attrs.items():
                node.set(qn(f"w:{key}"), str(value))

    @staticmethod
    def _add_rich_runs(paragraph, value):
        """Render the tiny safe markdown subset used by guide content."""
        text = str(value or "")
        position = 0
        for match in re.finditer(r"\*\*(.+?)\*\*", text):
            if match.start() > position:
                paragraph.add_run(text[position:match.start()])
            run = paragraph.add_run(match.group(1))
            run.bold = True
            position = match.end()
        if position < len(text):
            paragraph.add_run(text[position:])
        if not paragraph.runs:
            paragraph.add_run(text)
        return paragraph

    def _configure_styles(self, doc, guide=False):
        styles = doc.styles
        normal = styles["Normal"]
        normal.font.name = "Aptos"
        normal.font.size = Pt(10.5)
        normal.font.color.rgb = RGBColor(0x20, 0x2B, 0x33)
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.08

        title = styles["Title"]
        title.font.name = "Aptos Display"
        title.font.size = Pt(25 if guide else 24)
        title.font.bold = True
        title.font.color.rgb = RGBColor.from_string(self.NAVY if guide else "202B33")

        heading1 = styles["Heading 1"]
        heading1.font.name = "Aptos Display"
        heading1.font.size = Pt(16)
        heading1.font.bold = True
        heading1.font.color.rgb = RGBColor.from_string(self.TEAL if guide else "202B33")
        heading1.paragraph_format.space_before = Pt(14)
        heading1.paragraph_format.space_after = Pt(6)
        heading1.paragraph_format.keep_with_next = True

        heading2 = styles["Heading 2"]
        heading2.font.name = "Aptos Display"
        heading2.font.size = Pt(13)
        heading2.font.bold = True
        heading2.font.color.rgb = RGBColor.from_string(self.NAVY)

        for name in ["List Bullet", "List Number"]:
            styles[name].font.name = "Aptos"
            styles[name].font.size = Pt(10.5)

    def _configure_page(self, doc):
        for section in doc.sections:
            section.top_margin = Inches(0.72)
            section.bottom_margin = Inches(0.7)
            section.left_margin = Inches(0.8)
            section.right_margin = Inches(0.8)

    def _add_header_footer(self, doc, model, guide=False):
        for section in doc.sections:
            header = section.header
            p = header.paragraphs[0]
            p.text = (model.organisation or "Staff Guide") if guide else (model.organisation or "RazaAI")
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if guide else WD_ALIGN_PARAGRAPH.RIGHT
            for run in p.runs:
                run.font.name = "Aptos"
                run.font.size = Pt(8)
                run.font.bold = not guide
                run.font.color.rgb = RGBColor.from_string(self.MUTED if guide else "202B33")

            footer = section.footer
            fp = footer.paragraphs[0]
            if guide:
                fp.text = model.title
                fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
            else:
                left = model.document_id or "RazaAI Document"
                fp.text = f"{left}  |  Version {model.version}  |  {model.classification}"
                fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in fp.runs:
                run.font.name = "Aptos"
                run.font.size = Pt(8)
                run.font.color.rgb = RGBColor.from_string(self.MUTED if guide else "202B33")

    def _add_title_block(self, doc, model, guide=False):
        p = doc.add_paragraph()
        p.style = doc.styles["Title"]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(model.title)
        run.bold = True
        run.font.color.rgb = RGBColor.from_string(self.NAVY if guide else "202B33")

        if model.subtitle:
            p2 = doc.add_paragraph()
            p2.paragraph_format.space_after = Pt(10)
            r = p2.add_run(model.subtitle)
            r.font.name = "Aptos"
            r.font.size = Pt(12)
            r.font.color.rgb = RGBColor.from_string(self.MUTED)

        if guide:
            table = doc.add_table(rows=1, cols=1)
            cell = table.cell(0, 0)
            self._shade_cell(cell, self.TEAL)
            self._set_cell_margins(cell, top=20, bottom=20, start=0, end=0)
            cell.paragraphs[0].paragraph_format.space_after = Pt(0)
            doc.add_paragraph().paragraph_format.space_after = Pt(1)
            return

        created = model.created_date or date.today().isoformat()
        table = doc.add_table(rows=0, cols=2)
        items = [
            ("Document ID", model.document_id or "-"), ("Author", model.author),
            ("Status", model.status), ("Version", model.version),
            ("Classification", model.classification), ("Created", created),
        ]
        for key, value in items:
            row = table.add_row()
            row.cells[0].text = str(key); row.cells[1].text = str(value)
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                self._set_cell_margins(cell)
            for r in row.cells[0].paragraphs[0].runs:
                r.bold = True; r.font.size = Pt(9)
            for r in row.cells[1].paragraphs[0].runs:
                r.font.size = Pt(9)
        doc.add_paragraph()

    def _add_callout(self, doc, text):
        m = re.match(r"^(Tip|Important|Note|Example)\s*:\s*(.*)$", str(text).strip(), re.I)
        if not m:
            return False
        label, body = m.group(1).title(), m.group(2)
        bg = IMPORTANT_BG if label == "Important" else EXAMPLE_BG if label == "Example" else TIP_BG
        accent = "C99017" if label == "Important" else "4C8F5F" if label == "Example" else TEAL
        table = doc.add_table(rows=1, cols=1)
        cell = table.cell(0, 0)
        self._shade_cell(cell, bg)
        self._set_cell_margins(cell, top=120, start=150, bottom=120, end=150)
        self._set_cell_border(cell, left={"val":"single", "sz":"22", "color":accent})
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(f"{label}: "); r.bold = True
        self._add_rich_runs(p, body)
        doc.add_paragraph().paragraph_format.space_after = Pt(0)
        return True

    def _add_key_values(self, doc, items, guide=False):
        if not items:
            return
        table = doc.add_table(rows=0, cols=2)
        for item in items:
            row = table.add_row()
            row.cells[0].text = ""
            row.cells[1].text = ""
            self._add_rich_runs(row.cells[0].paragraphs[0], item.key)
            self._add_rich_runs(row.cells[1].paragraphs[0], item.value)
            if guide:
                self._shade_cell(row.cells[0], PALE_BLUE)
                self._set_cell_border(row.cells[0], left={"val":"single", "sz":"18", "color":self.TEAL})
            for cell in row.cells:
                self._set_cell_margins(cell, top=120 if guide else 90, start=120, bottom=120 if guide else 90, end=120)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for run in row.cells[0].paragraphs[0].runs:
                run.bold = True
                run.font.color.rgb = RGBColor.from_string(self.NAVY)
        doc.add_paragraph().paragraph_format.space_after = Pt(0)

    def _add_numbered_steps(self, doc, values, guide=False):
        if not values:
            return
        if not guide:
            for value in values:
                doc.add_paragraph(str(value), style="List Number")
            return
        for number, value in enumerate(values, 1):
            table = doc.add_table(rows=1, cols=2)
            table.columns[0].width = Inches(0.42)
            table.columns[1].width = Inches(6.35)
            num, body = table.rows[0].cells
            num.text = str(number)
            self._shade_cell(num, self.NAVY)
            self._shade_cell(body, PALE_GREY)
            for cell in (num, body):
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            self._set_cell_margins(num, top=115, start=90, bottom=115, end=90)
            self._set_cell_margins(body, top=115, start=150, bottom=115, end=150)
            for run in num.paragraphs[0].runs:
                run.bold = True; run.font.color.rgb = RGBColor(255,255,255); run.font.size = Pt(11)
            body.paragraphs[0].text = ""
            self._add_rich_runs(body.paragraphs[0], value)
            doc.add_paragraph().paragraph_format.space_after = Pt(0)

    def _add_table(self, doc, model_table, guide=False):
        table = doc.add_table(rows=1, cols=len(model_table.headers))
        table.style = "Table Grid"
        header = table.rows[0]
        self._set_repeat_table_header(header)
        for index, text in enumerate(model_table.headers):
            cell = header.cells[index]
            cell.text = ""
            self._add_rich_runs(cell.paragraphs[0], text)
            self._shade_cell(cell, self.NAVY if guide else "E7E6E6")
            for run in cell.paragraphs[0].runs:
                run.bold = True; run.font.size = Pt(9)
                if guide: run.font.color.rgb = RGBColor(255,255,255)
        for row_idx, values in enumerate(model_table.rows):
            row = table.add_row()
            for index, value in enumerate(values):
                cell = row.cells[index]
                cell.text = ""
                self._add_rich_runs(cell.paragraphs[0], value)
                if guide and row_idx % 2 == 1: self._shade_cell(cell, PALE_GREY)
                for run in cell.paragraphs[0].runs: run.font.size = Pt(9)
        for row in table.rows:
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
                self._set_cell_margins(cell, top=100, start=100, bottom=100, end=100)
        doc.add_paragraph()

    def write(self, model: DocumentModel, output_path: Path):
        model.validate()
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        guide = str(getattr(model, "style", "professional")).lower() == "guide"
        self._apply_theme(model)

        doc = Document()
        self._configure_styles(doc, guide=guide)
        self._configure_page(doc)
        self._add_header_footer(doc, model, guide=guide)
        self._add_title_block(doc, model, guide=guide)

        for section in model.sections:
            doc.add_heading(section.heading, level=1)
            for paragraph in section.paragraphs:
                if guide and self._add_callout(doc, paragraph):
                    continue
                p = doc.add_paragraph()
                self._add_rich_runs(p, paragraph)
            self._add_key_values(doc, section.key_values, guide=guide)
            self._add_numbered_steps(doc, section.numbered, guide=guide)
            for bullet in section.bullets:
                p = doc.add_paragraph(style="List Bullet")
                self._add_rich_runs(p, bullet)
            for table in section.tables:
                self._add_table(doc, table, guide=guide)

        properties = doc.core_properties
        properties.title = model.title
        properties.author = model.author
        properties.subject = model.subtitle or ""
        properties.comments = "Generated locally by RazaAI"
        doc.save(str(output_path))
        return output_path
