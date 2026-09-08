"""RazaAI staff rollout document completion and terminal execution."""

from pathlib import Path
import zipfile

from app.documents.guide_builder import (
    is_substantive_instruction_guide,
    normalize_instruction_guide_arguments,
)
from app.tools.documents import create_document


REQUEST = (
    "Create detailed docx for staff and the new FortiClient EMS filtering and "
    "protection being deployed throughout the school replacing Linewize."
)


def _docx_text(path):
    from docx import Document
    doc = Document(str(path))
    parts = []
    for p in doc.paragraphs:
        parts.append(p.text)
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def main():
    print("=" * 76)
    print("RazaAI Step 17.10 Staff Rollout Documents")
    print("=" * 76)

    sparse = {
        "title": "FortiClient EMS Deployment Guide",
        "format": "docx",
        "template": "instruction_guide",
        "subtitle": "Staff Guide",
        "sections": [
            {
                "heading": "Overview",
                "paragraphs": [
                    "The school is deploying FortiClient EMS as the new firewall "
                    "and network protection system to replace Linewize. FortiClient "
                    "EMS provides enhanced security, better application control, "
                    "and improved network management compared to Linewize."
                ],
            }
        ],
        "style": "guide",
    }

    normalized = normalize_instruction_guide_arguments(sparse, REQUEST)
    assert normalized["template"] == "instruction_guide"
    assert normalized["format"] == "docx"
    assert is_substantive_instruction_guide(normalized["sections"])
    assert normalized["title"] == "FortiClient EMS - Staff Deployment Guide"
    print("[PASS] sparse FortiClient rollout payload is expanded before execution")

    content = repr(normalized["sections"])
    assert "FortiClient EMS" in content
    assert "Linewize" in content
    assert "filtering and protection" in content.lower()
    assert "enhanced security" not in content.lower()
    assert "better application control" not in content.lower()
    assert "improved network management" not in content.lower()
    print("[PASS] supplied rollout facts survive while sparse-model inventions are discarded")

    headings = [section["heading"] for section in normalized["sections"]]
    for heading in (
        "What is changing",
        "What staff need to know",
        "What you may notice during rollout",
        "What staff should do",
        "If a website is blocked unexpectedly",
        "If FortiClient does not look right",
        "Privacy and passwords",
        "Frequently asked questions",
    ):
        assert heading in headings
    print("[PASS] detailed staff-facing rollout structure is present")

    result = create_document(
        **normalized,
        filename="step17_10_forticlient_staff_guide",
    )
    path = Path(result["path"])
    assert path.exists()
    text = _docx_text(path)
    assert "FortiClient EMS" in text
    assert "Linewize" in text
    assert "Frequently asked questions" in text
    assert "never include a password" in text.lower()
    print("[PASS] DOCX is created with substantive guide content")

    agent = Path("app/agent/agent.py").read_text(encoding="utf-8")
    assert "if tool_name == \"create_document\":" in agent
    assert "return final_document_message" in agent
    print("[PASS] create_document success/failure is terminal; no second Ollama wait")

    path.unlink(missing_ok=True)

    print()
    print("=" * 76)
    print("STEP 17.10 STAFF ROLLOUT DOCUMENTS PASSED")
    print("=" * 76)


if __name__ == "__main__":
    main()
