import tempfile
from pathlib import Path

from docx import Document

from app.documents import (
    DocumentEngine,
    DocumentModel,
    DocumentSection,
    DocumentTable,
    KeyValueItem,
)


def main():
    print()
    print("========================================")
    print("RazaAI Step 10.4 Document Engine Test")
    print("========================================")
    print()

    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)

        engine = DocumentEngine(
            project_root=root
        )

        model = DocumentModel(
            title="Clare Wi-Fi Incident Report",
            subtitle="Example-Staff Connectivity Failure",
            organisation="Example School",
            author="RazaAI",
            document_id="RAZA-INC-2026-0001",
            status="Final",
            classification="Internal",
            sections=[
                DocumentSection(
                    heading="Incident Summary",
                    paragraphs=[
                        (
                            "Staff wireless authentication succeeded, "
                            "but clients could not obtain usable "
                            "network connectivity."
                        )
                    ],
                    key_values=[
                        KeyValueItem(
                            "Site",
                            "Clare",
                        ),
                        KeyValueItem(
                            "SSID",
                            "Example-Staff",
                        ),
                    ],
                ),
                DocumentSection(
                    heading="Root Cause",
                    paragraphs=[
                        (
                            "The WLAN Access VLAN was configured as "
                            "VLAN 30 instead of VLAN 80."
                        )
                    ],
                ),
                DocumentSection(
                    heading="Resolution",
                    numbered=[
                        "Changed the WLAN Access VLAN from 30 to 80.",
                        "Reconnected the client.",
                    ],
                ),
                DocumentSection(
                    heading="Validation",
                    bullets=[
                        "Client received an IP address.",
                        "Gateway ping succeeded.",
                        "Internet connectivity was restored.",
                    ],
                ),
                DocumentSection(
                    heading="Evidence",
                    tables=[
                        DocumentTable(
                            headers=[
                                "Check",
                                "Result",
                            ],
                            rows=[
                                [
                                    "NPS authentication",
                                    "Passed",
                                ],
                                [
                                    "Access VLAN",
                                    "Corrected to 80",
                                ],
                                [
                                    "Internet",
                                    "Passed",
                                ],
                            ],
                        )
                    ],
                ),
            ],
        )

        result = engine.create_docx(
            model,
            filename="../../unsafe Clare report",
        )

        path = Path(
            result["path"]
        )

        if not path.exists():
            raise AssertionError(
                "DOCX was not created"
            )

        print("[PASS] DOCX generated")

        if (
            path.parent
            != root / "output" / "documents"
        ):
            raise AssertionError(
                "Document escaped safe output directory"
            )

        print(
            "[PASS] output path constrained to "
            "output/documents"
        )

        if ".." in path.name or "/" in path.name:
            raise AssertionError(
                "Unsafe filename was not sanitised"
            )

        print("[PASS] filename sanitised")

        doc = Document(
            str(path)
        )

        full_text = "\n".join(
            paragraph.text
            for paragraph in doc.paragraphs
        )

        for required in [
            "Clare Wi-Fi Incident Report",
            "Incident Summary",
            "Root Cause",
            "Resolution",
            "Validation",
            "Evidence",
        ]:
            if required not in full_text:
                raise AssertionError(
                    f"Missing generated text: "
                    f"{required}"
                )

        print("[PASS] document content verified")

        if len(doc.tables) < 3:
            raise AssertionError(
                "Expected metadata, key/value and evidence tables"
            )

        print("[PASS] structured tables generated")

        print()
        print(f"Generated sample: {path}")
        print()

    print("========================================")
    print("STEP 10.4 DOCUMENT ENGINE PASSED")
    print("========================================")
    print()


if __name__ == "__main__":
    main()
